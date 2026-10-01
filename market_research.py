"""Full extraction of the market research files (Strategy Plan .xlsx and the
HTML market reports) into three generic tables:

  market_reports   one row per source file: category, PIC, product line, plan date
  market_insights  every piece of written analysis (SWOT, conclusions, discount
                   patterns, customer insight, action plan text, report sections)
  market_tables    every data table, stored as column list + rows (JSON)

Used by ingest_v2.py (--market-dir). Only the standard library + openpyxl.
Cost / CM3 tables in the "2.Yes4All" sheets are skipped on purpose: the
dashboard database is readable with the public anon key.
"""
import html as htmllib
import os
import re
from datetime import date, datetime
from html.parser import HTMLParser

SECTION_RX = re.compile(r"^\s*\d+(\.\d+)*\.?\s+\S")
SENSITIVE_RX = re.compile(r"^(cost|cm3|fob|cogs|landed cost|margin)$", re.I)
META = {"pic:": "pic", "product line:": "product_line", "category:": "category_label", "plan date:": "plan_date"}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        return int(v) if v == int(v) and abs(v) < 1e12 else round(v, 4)
    if isinstance(v, str):
        s = v.replace("\r\n", "\n").replace("\r", "\n").strip()
        return s or None
    return v


def _is_text(v, n=1):
    return isinstance(v, str) and len(v) >= n


# ---------------------------------------------------------------------------
# Strategy Plan workbooks
# ---------------------------------------------------------------------------
def xlsx_research(path, cat):
    import openpyxl
    import warnings
    warnings.filterwarnings("ignore", category=UserWarning)
    src = os.path.basename(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    report = {"category": cat, "source": src, "kind": "Strategy Plan (xlsx)", "pic": None, "product_line": None, "plan_date": None, "title": None}
    insights, tables = [], []
    for ws in wb.worksheets:
        if not re.match(r"^\s*\d", ws.title):
            continue
        sheet = ws.title.strip()
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        ex_col = None
        for r in raw[:4]:
            for j, c in enumerate(r):
                if isinstance(c, str) and c.strip() == "Ex":
                    ex_col = j
        rows = []
        for r in raw:
            r = r[:ex_col] if ex_col is not None else r
            rows.append([(j, _clean(c)) for j, c in enumerate(r) if _clean(c) not in (None, "")])
        section, label, skip = sheet, None, False
        i = 0
        while i < len(rows):
            cells = rows[i]
            if not cells:
                i += 1
                continue
            first = cells[0][1]
            # header meta (PIC / Product line / Category / Plan date)
            if _is_text(first) and first.strip().lower() in META or any(_is_text(v) and v.strip().lower() in META for _, v in cells[:2]):
                for k, (j, v) in enumerate(cells):
                    key = META.get(str(v).strip().lower()) if _is_text(v) else None
                    if key and k + 1 < len(cells) and not report.get(key):
                        report[key] = cells[k + 1][1]
                i += 1
                continue
            if _is_text(first) and first.strip().upper().startswith("STRATEGY PLAN") and len(cells) == 1:
                report["title"] = report["title"] or first
                i += 1
                continue
            if _is_text(first) and SECTION_RX.match(first) and len(first) < 90:
                section, label, skip = first.strip(), None, False
                i += 1
                continue
            if _is_text(first) and first.strip().lower() in ("ex:", "ex"):
                skip = True  # template guidance until the next section
            if skip:
                i += 1
                continue
            # contiguous block
            block = []
            while i < len(rows) and rows[i]:
                c0 = rows[i][0][1]
                if _is_text(c0) and SECTION_RX.match(c0) and len(c0) < 90:
                    break
                block.append(rows[i])
                i += 1
            # leading single short cells are labels / titles
            while block and len(block) > 1 and len(block[0]) <= 2 and all(_is_text(v) and len(v) < 80 for _, v in block[0]) \
                    and (len(block[0]) == 1 or len(block[1]) > len(block[0]) + 1):
                label = " · ".join(v for _, v in block.pop(0))
            if not block:
                continue
            multi = [r for r in block if len(r) >= 3]
            if len(block) >= 2 and multi:
                # a lone short label inside the block starts a new table ("Key action & Timeline")
                part, plabel = [], label
                for r in block:
                    if len(r) == 1 and _is_text(r[0][1]) and len(r[0][1]) < 80 and "\n" not in r[0][1] and part:
                        if len(part) >= 2:
                            _add_table(tables, cat, src, sheet, section, plabel, part)
                        part, plabel = [], r[0][1]
                        continue
                    part.append(r)
                if len(part) >= 2:
                    _add_table(tables, cat, src, sheet, section, plabel, part)
                label = plabel
                continue
            for r in block:
                if len(r) == 1:
                    v = r[0][1]
                    if _is_text(v) and len(v) < 60 and "\n" not in v:
                        label = v
                    elif v is not None:
                        insights.append(_ins(cat, src, sheet, section, label, str(v)))
                else:
                    head = r[0][1]
                    rest = "\n".join(str(v) for _, v in r[1:])
                    if _is_text(head) and len(str(head)) < 80 and len(rest) > 40:
                        insights.append(_ins(cat, src, sheet, section, head.rstrip(":"), rest))
                    elif len(" ".join(str(v) for _, v in r)) > 60:
                        insights.append(_ins(cat, src, sheet, section, label, " · ".join(str(v) for _, v in r)))
    return report, insights, tables


def _ins(cat, src, sheet, section, label, content):
    return {"category": cat, "source": src, "sheet": sheet, "section": section, "label": (label or "").rstrip(":").strip() or None,
            "content": content.strip()[:8000]}


def _add_table(tables, cat, src, sheet, section, label, block):
    # optional group header ("Cost | CM3") above the real header
    groups = None
    if len(block) >= 3 and len(block[0]) < len(block[1]) and all(_is_text(v) for _, v in block[0]) and len(block[0]) <= 4:
        groups = block.pop(0)
    head = block[0]
    if any(SENSITIVE_RX.match(str(v).strip()) for _, v in (groups or []) + head) or re.search(r"\b(cost|cm3|fob)\b", label or "", re.I):
        return
    cols = sorted({j for r in block for j, _ in r})
    hdr = dict(head)
    header_is_text = sum(1 for _, v in head if _is_text(v)) >= max(1, len(head) // 2) or len(head) >= len(block[1]) if len(block) > 1 else True
    if header_is_text:
        names = [str(hdr.get(j, "")) or f"col{k + 1}" for k, j in enumerate(cols)]
        body = block[1:]
    else:
        names = [f"col{k + 1}" for k in range(len(cols))]
        body = block
    seen = {}
    for k, n in enumerate(names):
        if n in seen:
            seen[n] += 1
            names[k] = f"{n} ({seen[n]})"
        else:
            seen[n] = 1
    data = []
    for r in body:
        d = dict(r)
        data.append([d.get(j) for j in cols])
    if not data:
        return
    # drop columns that are empty in every data row and unnamed
    keep = [k for k in range(len(cols)) if any(row[k] is not None for row in data) or not names[k].startswith("col")]
    names = [names[k] for k in keep]
    data = [[row[k] for k in keep] for row in data]
    tables.append({"category": cat, "source": src, "sheet": sheet, "section": section, "title": label or section,
                   "columns": names, "rows": data})


# ---------------------------------------------------------------------------
# HTML reports
# ---------------------------------------------------------------------------
class _ReportParser(HTMLParser):
    BLOCK = {"p", "li", "div", "br", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "blockquote", "ul", "ol"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.section = None
        self.head_tag = None
        self.head_buf = []
        self.text = {}      # section -> list of lines
        self.order = []
        self.line = []
        self.tables = []
        self.tstack = []    # open tables: {"rows": [], "row": None, "cell": None}
        self.title = None
        self.in_title = False

    def _flush(self):
        s = re.sub(r"\s+", " ", "".join(self.line)).strip()
        self.line = []
        if s and len(s) > 1:
            sec = self.section or "Tổng quan"
            if sec not in self.text:
                self.text[sec] = []
                self.order.append(sec)
            if not self.text[sec] or self.text[sec][-1] != s:
                self.text[sec].append(s)

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "svg", "noscript", "template"):
            self.skip += 1
            return
        if self.skip:
            return
        if tag == "title":
            self.in_title = True
        if tag in ("h1", "h2", "h3", "h4"):
            self._flush()
            self.head_tag, self.head_buf = tag, []
            return
        if tag == "table":
            self._flush()
            self.tstack.append({"rows": [], "row": None, "cell": None, "section": self.section})
            return
        if self.tstack:
            t = self.tstack[-1]
            if tag == "tr":
                t["row"] = []
            elif tag in ("td", "th"):
                t["cell"] = []
                span = dict(attrs).get("colspan")
                t["span"] = int(span) if span and str(span).isdigit() else 1
            elif tag == "br" and t["cell"] is not None:
                t["cell"].append(" ")
            return
        if tag in self.BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in ("script", "style", "svg", "noscript", "template"):
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag == "title":
            self.in_title = False
        if self.head_tag and tag == self.head_tag:
            h = re.sub(r"\s+", " ", "".join(self.head_buf)).strip()
            if h:
                self.section = h[:120]
            self.head_tag = None
            return
        if self.tstack:
            t = self.tstack[-1]
            if tag in ("td", "th") and t["cell"] is not None:
                v = re.sub(r"\s+", " ", "".join(t["cell"])).strip()
                t["row"] = (t["row"] or []) + [v] + [""] * (t.get("span", 1) - 1)
                t["cell"] = None
            elif tag == "tr" and t["row"] is not None:
                if any(t["row"]):
                    t["rows"].append(t["row"])
                t["row"] = None
            elif tag == "table":
                done = self.tstack.pop()
                if len(done["rows"]) >= 2:
                    self.tables.append((done["section"] or self.section, done["rows"]))
            return
        if tag in self.BLOCK:
            self._flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.in_title:
            self.title = ((self.title or "") + data).strip()
        if self.head_tag:
            self.head_buf.append(data)
            return
        if self.tstack:
            t = self.tstack[-1]
            if t["cell"] is not None:
                t["cell"].append(data)
            return
        self.line.append(data)


def _num(s):
    if not isinstance(s, str):
        return s
    t = s.strip().replace(",", "")
    m = re.fullmatch(r"\$?(-?\d+(?:\.\d+)?)([kKmM%]?)", t)
    if not m:
        return s.strip() or None
    v = float(m.group(1))
    suf = m.group(2).lower()
    if suf == "k":
        v *= 1e3
    elif suf == "m":
        v *= 1e6
    elif suf == "%":
        v /= 100
    return int(v) if v == int(v) else round(v, 4)


# ---------------------------------------------------------------------------
# Reports that embed their data as JSON (<script id="report-data">)
# ---------------------------------------------------------------------------
DROP_KEY_RX = re.compile(r"cm3|cost|margin|fob|slug|^url$|is_estimate|elasticity_is_default", re.I)
LABELS = {"exec_summary": "Tóm tắt cho người ra quyết định", "market": "Tổng quan thị trường", "market_bubble": "Quy mô thị trường theo dòng sản phẩm",
          "market_category_trend": "Diễn biến giá & số bán theo dòng sản phẩm", "inventory_table": "Tồn kho & cần chú ý", "rollup": "Tổng hợp đề xuất & mục tiêu",
          "sku_master": "Bảng giá đề xuất — tất cả SKU", "strategy_by_line": "Chiến lược theo dòng sản phẩm", "rows": "Giá Yes4All vs đối thủ",
          "competitor_overview": "Đối thủ đã đối chiếu", "price_reco": "Đề xuất giá", "callouts": "Ghi chú SKU", "corr_market": "Giá & số bán theo brand (tháng)",
          "inv_corr": "Tồn kho vs số bán theo SKU", "units_corr": "Giá vs số bán theo SKU", "kpis": "KPI"}


def _txt(v):
    return htmllib.unescape(re.sub(r"<[^>]+>", "", v)).strip() if isinstance(v, str) else v


def _scalar(v):
    if isinstance(v, (dict, list)):
        if isinstance(v, list) and all(not isinstance(x, (dict, list)) for x in v):
            return ", ".join(str(x) for x in v if x is not None)[:500] or None
        if isinstance(v, list) and all(isinstance(x, dict) for x in v):
            return "; ".join(" ".join(str(y) for k, y in x.items() if not DROP_KEY_RX.search(k) and not isinstance(y, (dict, list))) for x in v)[:500]
        if isinstance(v, dict):
            return ", ".join(f"{k}={y}" for k, y in v.items() if not DROP_KEY_RX.search(k) and not isinstance(y, (dict, list)))[:500]
        return None
    if isinstance(v, float):
        return round(v, 4)
    return _txt(v)


def _json_walk(o, key, section, cat, src, insights, tables):
    title = LABELS.get(key, key.replace("_", " ").capitalize() if key else section)
    if isinstance(o, str):
        if len(o) > 30:
            insights.append({"category": cat, "source": src, "sheet": "report", "section": section, "label": title, "content": _txt(o)[:8000]})
        return
    if isinstance(o, list):
        if o and all(isinstance(x, str) for x in o):
            if sum(len(x) for x in o) / len(o) < 25:
                return
            insights.append({"category": cat, "source": src, "sheet": "report", "section": section, "label": title, "content": "\n".join("- " + _txt(x) for x in o)[:8000]})
            return
        if o and all(isinstance(x, dict) for x in o) and any(isinstance(v, list) and v and isinstance(v[0], dict) for x in o for v in x.values()) \
                and not all(isinstance(x.get("months"), list) for x in o):
            for x in o:  # complex records (one product line each): walk them one by one
                _json_walk(x, None, x.get("name") or x.get("pl") or section, cat, src, insights, tables)
            return
        if o and all(isinstance(x, dict) for x in o):
            # rows of parallel arrays (months + values) become long tables
            if all(isinstance(x.get("months"), list) for x in o):
                cols, rows = None, []
                for x in o:
                    arrays = {k: v for k, v in x.items() if isinstance(v, list) and len(v) == len(x["months"]) and k != "months"}
                    ident = {k: v for k, v in x.items() if not isinstance(v, (list, dict)) and not DROP_KEY_RX.search(k)}
                    cols = list(ident) + ["month"] + list(arrays)
                    for i, m in enumerate(x["months"]):
                        rows.append(list(ident.values()) + [m] + [_scalar(arrays[k][i]) for k in arrays])
                if rows:
                    tables.append({"category": cat, "source": src, "sheet": "report", "section": section, "title": title, "columns": cols, "rows": rows})
                return
            keys = []
            for x in o:
                for k, v in x.items():
                    if k not in keys and not DROP_KEY_RX.search(k):
                        keys.append(k)
            rows = [[_scalar(x.get(k)) for k in keys] for x in o]
            keys2 = [k for i, k in enumerate(keys) if any(r[i] not in (None, "") for r in rows)]
            idx = [keys.index(k) for k in keys2]
            tables.append({"category": cat, "source": src, "sheet": "report", "section": section, "title": title,
                           "columns": keys2, "rows": [[r[i] for i in idx] for r in rows]})
            # long text fields inside rows (reasoning, notes) are also kept as insights
            for x in o:
                name = x.get("sku") or x.get("pl") or x.get("brand") or ""
                for k, v in x.items():
                    if isinstance(v, str) and len(v) > 120 and not DROP_KEY_RX.search(k):
                        insights.append({"category": cat, "source": src, "sheet": "report", "section": section, "label": f"{title} · {name}".strip(" ·"), "content": _txt(v)[:8000]})
            return
        return
    if isinstance(o, dict):
        sub_section = o.get("name") or o.get("pl") if isinstance(o.get("name") or o.get("pl"), str) else None
        sec = sub_section or (section if key in LABELS and key not in ('market', 'market_bubble', 'market_category_trend', 'inventory_table', 'rollup') else (title if key else section))
        months = o.get("months") if isinstance(o.get("months"), list) else None
        flat = []
        for k, v in o.items():
            if DROP_KEY_RX.search(k):
                continue
            if months is not None and k != "months" and isinstance(v, (list, dict)):
                continue
            if isinstance(v, (dict, list)):
                _json_walk(v, k, sec, cat, src, insights, tables)
            elif isinstance(v, str) and len(v) > 30:
                insights.append({"category": cat, "source": src, "sheet": "report", "section": sec, "label": LABELS.get(k, k), "content": _txt(v)[:8000]})
            elif v is not None and k not in ("name", "pl"):
                flat.append(f"{k}: {_scalar(v)}")
        if months is not None:
            series = {}
            for k, v in o.items():
                if k == "months" or DROP_KEY_RX.search(k):
                    continue
                if isinstance(v, list) and len(v) == len(months) and all(not isinstance(x, (dict, list)) for x in v):
                    series[k] = v
                elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                    for x in v:
                        name = x.get("brand") or x.get("pl") or x.get("sku") or ""
                        for kk, vv in x.items():
                            if isinstance(vv, list) and len(vv) == len(months):
                                series[f"{name} · {kk}"] = vv
                elif isinstance(v, dict):
                    for kk, vv in v.items():
                        if isinstance(vv, list) and len(vv) == len(months):
                            series[f"{k} · {kk}"] = vv
            if series:
                tables.append({"category": cat, "source": src, "sheet": "report", "section": sec, "title": title, "columns": ["month"] + list(series),
                               "rows": [[m] + [_scalar(series[k][i]) for k in series] for i, m in enumerate(months)]})
        if flat and len(flat) >= 2 and key:
            insights.append({"category": cat, "source": src, "sheet": "report", "section": sec, "label": title + " — số liệu", "content": "\n".join(flat)[:8000]})


def json_report(raw, cat, src, insights, tables):
    m = re.search(r'<script[^>]*id="report-data"[^>]*>(.*?)</script>', raw, re.S)
    if not m:
        return False
    import json
    try:
        d = json.loads(m.group(1))
    except ValueError:
        return False
    for k, v in d.items():
        if k == "meta":
            continue
        _json_walk(v, k, LABELS.get(k, k), cat, src, insights, tables)
    return True


def html_research(path, cat):
    src = os.path.basename(path)
    raw = open(path, encoding="utf-8", errors="ignore").read()
    p = _ReportParser()
    p.feed(raw)
    p._flush()
    title = htmllib.unescape(p.title or os.path.splitext(src)[0])
    report = {"category": cat, "source": src, "kind": "Market report (html)", "pic": None, "product_line": None, "plan_date": None, "title": title}
    insights, tables = [], []
    for sec in p.order:
        lines = [l for l in p.text[sec] if len(l) > 2]
        body = "\n".join(lines).strip()
        if len(body) >= 80:  # shorter blocks are menus / page titles
            for k in range(0, len(body), 8000):
                insights.append({"category": cat, "source": src, "sheet": "report", "section": sec, "label": None, "content": body[k:k + 8000]})
    json_report(raw, cat, src, insights, tables)
    for sec, rows in p.tables:
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        head, body = rows[0], rows[1:]
        names, seen = [], {}
        for k, n in enumerate(head):
            n = n or f"col{k + 1}"
            seen[n] = seen.get(n, 0) + 1
            names.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
        if any(SENSITIVE_RX.match(n.strip()) for n in names):
            continue
        tables.append({"category": cat, "source": src, "sheet": "report", "section": sec or title, "title": sec or title,
                       "columns": names, "rows": [[_num(v) for v in r] for r in body]})
    return report, insights, tables


def parse_research(market_dir, category_of):
    import glob
    reports, insights, tables = [], [], []
    for path in sorted(glob.glob(os.path.join(market_dir, "*.xlsx")) + glob.glob(os.path.join(market_dir, "*.html"))):
        cat = category_of(os.path.basename(path))
        try:
            r, i, t = (xlsx_research if path.endswith(".xlsx") else html_research)(path, cat)
        except Exception as e:  # keep going: one bad file should not drop the rest
            print(f"  ! {os.path.basename(path)}: {e}")
            continue
        reports.append(r)
        insights.extend(i)
        tables.extend(t)
    tables = [t for t in (_drop_cost_columns(t) for t in tables) if t]
    for k, x in enumerate(insights):
        x["sort"] = k
    for k, x in enumerate(tables):
        x["sort"] = k
    return reports, insights, tables


COST_COL_RX = re.compile(r"cm3|fob|cogs|^cost\b|\bcost (avc|asc)", re.I)


def _drop_cost_columns(t):
    """Remove CM3 / cost columns (and the AVC sell-in price columns that sit
    under the Cost group) so no cost data reaches the public-readable tables."""
    cols = [str(c) for c in t["columns"]]
    if not any(COST_COL_RX.search(c) for c in cols):
        return t
    drop = {i for i, c in enumerate(cols) if COST_COL_RX.search(c) or re.fullmatch(r"AVC (DS|WH|DI)( \(\d+\))?", c.strip())}
    keep = [i for i in range(len(cols)) if i not in drop]
    if len(keep) < 2:
        return None
    t["columns"] = [t["columns"][i] for i in keep]
    t["rows"] = [[r[i] if i < len(r) else None for i in keep] for r in t["rows"]]
    return t


# ---------------------------------------------------------------------------
# Market / competitor focus: what the Market tab shows
# ---------------------------------------------------------------------------
KEEP_RX = re.compile(r"brand|đối thủ|competitor|market|thị trường|keyword|buy box|variation|channel|top brands|category mapping|parent asin", re.I)
DROP_RX = re.compile(r"by sku|theo sku|npd|action plan|inventory|stock|tồn kho|đề xuất giá|kpi|content|yoy|monthly performance|performance by|"
                     r"ghi chú|phương pháp|customer insight|timeline|real sale|doanh số theo sku|dead stock|review needed|chiến lược", re.I)


def is_market_table(t):
    title = " ".join(str(x) for x in (t.get("title"), t.get("section")) if x)
    cols = " ".join(str(c) for c in t["columns"])
    if t.get("sheet") not in (None, "report", "1.Market"):
        return False  # 2.Yes4All / 3.Action plan sheets are about our own SKUs
    if DROP_RX.search(title) and not re.search(r"đối thủ|competitor", title, re.I):
        return False
    return bool(KEEP_RX.search(title) or KEEP_RX.search(cols))


def balance_board_lines(raw):
    """Product-line names in the Balance Board report (each becomes a category)."""
    m = re.search(r'<script[^>]*id="report-data"[^>]*>(.*?)</script>', raw, re.S)
    if not m:
        return []
    import json
    try:
        return [p.get("name") for p in json.loads(m.group(1)).get("product_lines", []) if p.get("name")]
    except ValueError:
        return []


def market_library(market_dir, category_of):
    """Market/competitor tables only, for the Market tab and the AI chat."""
    import glob
    reports, _insights, tables = parse_research(market_dir, category_of)
    lines = {}
    for path in glob.glob(os.path.join(market_dir, "*.html")):
        for n in balance_board_lines(open(path, encoding="utf-8", errors="ignore").read()):
            lines[n] = os.path.basename(path)
    out = []
    for t in tables:
        if not is_market_table(t):
            continue
        if t["source"] in lines.values() and t.get("section") in lines:
            t["category"] = t["section"]  # Balance Board: one category per product line
        out.append(t)
    for n, src in lines.items():
        reports.append({"category": n, "source": f"{src}#{n}", "kind": "Market report (html)", "title": f"Balance Board — {n}",
                        "pic": None, "product_line": n, "plan_date": None})
    for k, x in enumerate(out):
        x["sort"] = k
    return reports, out


def json_brand_monthly(raw, src):
    """Balance Board report: monthly price & units per competitor brand, per product line."""
    m = re.search(r'<script[^>]*id="report-data"[^>]*>(.*?)</script>', raw, re.S)
    if not m:
        return []
    import json
    d = json.loads(m.group(1))
    rows = []
    for p in d.get("product_lines", []):
        cm = p.get("corr_market") or {}
        months = cm.get("months") or []
        series = [(b.get("brand"), b.get("price") or [], b.get("units") or []) for b in cm.get("brands", [])]
        y = cm.get("yes4all") or {}
        if y:
            series.append(("Yes4All", y.get("price") or [], y.get("units") or []))
        for brand, pr, un in series:
            for i, mo in enumerate(months):
                pv = pr[i] if i < len(pr) else None
                uv = un[i] if i < len(un) else None
                if not uv and not pv:
                    continue
                mo_d = str(mo)[:7] + "-01"
                rev = round(pv * uv, 2) if pv and uv else None
                rows.append({"category": p.get("name"), "month": mo_d, "brand": brand, "revenue": rev if rev is not None else 0,
                             "units": uv, "avg_price": pv, "source": f"{src}#{p.get('name')}"})
    return rows


def variation_from_tables(tables, period="Apr-Jun 2026"):
    """Brand × variation tables in the HTML reports ("Market Size by Brand & Color", "Price by Brand & Size")."""
    out = []
    skip = re.compile(r"asin|link|total|share|rating|rank|discount|pattern|note|channel|market size \(|^col\d|avg_|units|price", re.I)
    for t in tables:
        cols = [str(c) for c in t["columns"]]
        if not cols or not re.match(r"^\s*brand", cols[0], re.I) or t.get("sheet") not in (None, "report"):
            continue
        title = f"{t.get('title') or ''} {cols[0]}"
        if not re.search(r"market size|units|price|asp|giá|doanh thu|revenue", title, re.I):
            continue  # cannot tell whether the cells are sizes or prices
        metric = "price" if re.search(r"price|asp|giá", title, re.I) else "revenue"
        var_idx = [(k, c) for k, c in enumerate(cols) if k > 0 and not skip.search(c)]
        if len(var_idx) < 2:
            continue
        for r in t["rows"]:
            name = r[0]
            if not isinstance(name, str) or not name.strip() or re.match(r"^(total|min|median|max|y4a vs)", name.strip(), re.I):
                continue
            for k, var in var_idx:
                v = r[k] if k < len(r) else None
                if isinstance(v, str):
                    v = _num(v)
                if not isinstance(v, (int, float)) or v <= 0:
                    continue
                out.append({"category": t["category"], "attribute": "variation", "variation": var, "brand": name.strip(),
                            "metric": metric, "value": round(float(v), 2), "period": period, "source": t["source"]})
    return out
