"""
Yes4All Sales Dashboard — Supabase ingestion script.

Reads the source files (follow-up Excel, target HTML, daily/hourly sales
Excel) and prepares the matching upsert into Supabase. Two modes:

  sql (default) — writes a .sql file with INSERT ... ON CONFLICT statements
                  that you paste into the Supabase SQL Editor yourself and
                  run. Use this when ingest.py is run from an environment
                  without direct internet access to your Supabase project
                  (e.g. inside a Cowork/Claude sandbox) — this is the mode
                  Claude will normally use.

  api            — pushes straight to Supabase's REST API over the network.
                  Only works where the machine running this script has real
                  internet access to *.supabase.co (i.e. run it yourself,
                  locally, not from a sandboxed agent). Needs:
                    SUPABASE_URL          e.g. https://xxxxxxxx.supabase.co
                    SUPABASE_SERVICE_KEY  the "service_role" key (Project
                                          Settings -> API) — NOT the anon key.
                  Select it with: export INGEST_MODE=api

Usage:
    python ingest.py followup "Amazon Follow up_2026.xlsx"
    python ingest.py target   "SSO_US_Aug.2026.html" --month 2026-08-01
    python ingest.py sales    "SSO Data Extraction Hourly -- hourly.xlsx"

    # optional, only for --mode api:
    export SUPABASE_URL=https://xxxxxxxx.supabase.co
    export SUPABASE_SERVICE_KEY=eyJ...

Every statement is an upsert keyed on the table's natural key (SKU, or
SKU+date, or SKU+month), so re-running with the same or overlapping data is
always safe — it overwrites matching rows instead of duplicating them.
"""
import os
import re
import sys
import csv
import json
import argparse
from datetime import datetime

import requests
import openpyxl

INGEST_MODE = os.environ.get("INGEST_MODE", "sql")
SQL_BUFFER = []


# ---------------------------------------------------------------------------
# Supabase REST helpers
# ---------------------------------------------------------------------------
def _config():
    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit(
            "Missing SUPABASE_URL / SUPABASE_SERVICE_KEY environment variables.\n"
            "Set them first, e.g.:\n"
            "  export SUPABASE_URL=https://xxxxxxxx.supabase.co\n"
            "  export SUPABASE_SERVICE_KEY=eyJ...\n"
        )
    return url.rstrip("/"), key


def sql_literal(v):
    """Render a Python value as a Postgres SQL literal."""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return repr(v)
    s = str(v).replace("'", "''")
    return f"'{s}'"


def safe_num(v, default=0):
    """Coerce an Excel cell value to a number, defaulting to 0 for blanks or
    Excel error strings (#N/A, #REF!, #VALUE! etc.) instead of letting them
    leak through and break the numeric column in Postgres."""
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return v
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def rows_to_sql(table, rows, on_conflict, batch_size=500):
    """Build INSERT ... ON CONFLICT DO UPDATE statements for `rows`, batched."""
    if not rows:
        return ""
    cols = list(rows[0].keys())
    conflict_cols = [c.strip() for c in on_conflict.split(",")]
    update_cols = [c for c in cols if c not in conflict_cols]
    set_clause = ", ".join(f"{c}=excluded.{c}" for c in update_cols)
    statements = []
    for i in range(0, len(rows), batch_size):
        chunk = rows[i:i + batch_size]
        values_sql = ",\n".join(
            "  (" + ", ".join(sql_literal(r.get(c)) for c in cols) + ")" for r in chunk
        )
        statements.append(
            f"INSERT INTO {table} ({', '.join(cols)}) VALUES\n{values_sql}\n"
            f"ON CONFLICT ({', '.join(conflict_cols)}) DO UPDATE SET {set_clause};"
        )
    return "\n\n".join(statements)


def upsert_rows_api(table, rows, on_conflict, batch_size=500):
    """Push straight to Supabase's REST API. Only works with real internet
    access to *.supabase.co — see INGEST_MODE in the module docstring."""
    if not rows:
        print(f"[{table}] nothing to upsert")
        return
    url, key = _config()
    endpoint = f"{url}/rest/v1/{table}?on_conflict={on_conflict}"
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
    }
    total = len(rows)
    for i in range(0, total, batch_size):
        chunk = rows[i:i + batch_size]
        resp = requests.post(endpoint, headers=headers, data=json.dumps(chunk, default=str), timeout=60)
        if resp.status_code >= 300:
            raise RuntimeError(f"[{table}] upsert failed ({resp.status_code}): {resp.text[:500]}")
        print(f"[{table}] upserted rows {i+1}-{min(i+batch_size, total)} / {total}")
    print(f"[{table}] done — {total} rows upserted (api)")


def upsert_rows(table, rows, on_conflict, batch_size=500):
    """Upsert a list of dict rows into `table`, keyed on `on_conflict`
    (comma-separated column names). Mode controlled by INGEST_MODE env var:
    'sql' (default) buffers SQL to be written to a .sql file by main();
    'api' pushes straight to Supabase over the network (needs real internet)."""
    if not rows:
        print(f"[{table}] nothing to upsert")
        return
    if INGEST_MODE == "api":
        upsert_rows_api(table, rows, on_conflict, batch_size)
        return
    sql = rows_to_sql(table, rows, on_conflict, batch_size)
    SQL_BUFFER.append(f"-- {table}: {len(rows)} row(s)\n{sql}")
    print(f"[{table}] prepared {len(rows)} row(s) of SQL (mode=sql)")


# ---------------------------------------------------------------------------
# 1) Follow-up file -> skus (PIC / Main PL / Sub PL / inventory / identity)
# ---------------------------------------------------------------------------
def load_managed_skus(path, sheet="Tracking (2)"):
    """Read just the SKU column (B) from the follow-up/tracking file — this is
    the authoritative list of SKUs Rachel's team actually manages. Pass the
    result as sku_filter to ingest_target/ingest_sales_excel so the company-
    wide SSO planning/sales exports don't pollute the database with hundreds
    of SKUs outside her scope (which show up as PIC=Unassigned / mainPL/subPL
    =Unclassified placeholders in the dashboard)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    skus = set()
    for r in range(5, ws.max_row + 1):
        sku = ws.cell(row=r, column=2).value
        if sku:
            skus.add(str(sku).strip())
    wb.close()
    return skus


def sql_in_list(values):
    return ", ".join(sql_literal(v) for v in values)


def cleanup_unmanaged_skus(followup_path):
    """One-time fix: DELETE any `skus` row whose SKU is NOT in the follow-up
    file's managed list. Needed because upserts only ever add/update rows —
    they never remove the ~1000 unrelated-catalog placeholder SKU rows that
    got ingested before --skus-from existed. targets_monthly/sales_daily rows
    for those SKUs cascade-delete automatically (FK ... ON DELETE CASCADE);
    diary_entries/projects are untouched (they key off main_pl/project name,
    not sku)."""
    managed = load_managed_skus(followup_path)
    if not managed:
        print("[cleanup] no managed SKUs found in the follow-up file — aborting, would delete everything")
        return
    stmt = f"DELETE FROM skus WHERE sku NOT IN ({sql_in_list(sorted(managed))});"
    SQL_BUFFER.append(
        f"-- cleanup: remove any skus row NOT in the {len(managed)} managed SKU(s) from the follow-up file\n"
        f"-- (targets_monthly / sales_daily rows for those SKUs cascade-delete automatically)\n{stmt}"
    )
    print(f"[cleanup] prepared DELETE for skus NOT IN {len(managed)} managed SKU(s)")


def ingest_followup(path, sheet="Tracking (2)"):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = {}
    for r in range(5, ws.max_row + 1):
        sku = ws.cell(row=r, column=2).value
        if not sku:
            continue
        sku = str(sku).strip()
        if sku in rows:
            continue  # same SKU repeated under multiple ASINs — identity fields are identical
        rows[sku] = {
            "sku": sku,
            "product_name": ws.cell(row=r, column=5).value,
            "pic": ws.cell(row=r, column=1).value or "Unassigned",
            "sub_pl": ws.cell(row=r, column=6).value or "Unclassified",
            "main_pl": ws.cell(row=r, column=7).value or "Unclassified",
            "lifecycle": ws.cell(row=r, column=8).value,
            "selling_type": ws.cell(row=r, column=9).value,
            "asin_status": ws.cell(row=r, column=10).value,
            "salable_y4a": safe_num(ws.cell(row=r, column=11).value),   # column K
            "salable_amz": safe_num(ws.cell(row=r, column=12).value),   # column L
            "updated_at": datetime.utcnow().isoformat(),
        }
    wb.close()
    upsert_rows("skus", list(rows.values()), on_conflict="sku")


# ---------------------------------------------------------------------------
# 2) Target file (SSO planning HTML) -> skus (channel/portfolio/moc/rrp) +
#    targets_monthly (units/GMV/ads/promo target for one calendar month)
# ---------------------------------------------------------------------------
MONTH_LABEL_RE = re.compile(r"([A-Za-z]{3})-(\d{2})")
MONTHS = {"Jan":1,"Feb":2,"Mar":3,"Apr":4,"May":5,"Jun":6,"Jul":7,"Aug":8,"Sep":9,"Oct":10,"Nov":11,"Dec":12}


def _label_to_date(label):
    m = MONTH_LABEL_RE.match(label)
    if not m:
        raise ValueError(f"Cannot parse month label: {label}")
    mon, yy = m.group(1), int(m.group(2))
    year = 2000 + yy
    return f"{year:04d}-{MONTHS[mon]:02d}-01"


def ingest_target(path, month_label=None, month_index=0, sku_filter=None):
    """month_label e.g. '2026-08-01' (already a date) or 'Aug-26' (SSO label).
    month_index selects which entry of the payload's futureLabels/final[] array
    to use if you want a month other than the first (0 = first future month).
    sku_filter: optional set of SKUs — if given, only those SKUs are ingested
    (the SSO planning export covers the whole company's catalog, not just the
    SKUs Rachel's team manages)."""
    with open(path, encoding="utf-8", errors="ignore") as f:
        content = f.read()
    m = re.search(r'<script id="payload" type="application/json">(.*?)</script>', content, re.S)
    payload = json.loads(m.group(1))
    schema = payload["schema"]
    idx = {fld: i for i, fld in enumerate(schema)}
    strings = payload["strings"]
    string_fields = set(payload["stringFields"])
    future_labels = payload["meta"]["futureLabels"]

    if month_label is None:
        month_date = _label_to_date(future_labels[month_index])
    elif re.match(r"^\d{4}-\d{2}-\d{2}$", month_label):
        month_date = month_label
    else:
        month_date = _label_to_date(month_label)

    def resolve(row, field):
        v = row[idx[field]]
        return strings[v] if field in string_fields and v is not None else v

    sku_updates, target_rows = [], []
    skipped = 0
    for row in payload["rows"]:
        sku = resolve(row, "sku")
        if sku_filter is not None and sku not in sku_filter:
            skipped += 1
            continue
        final_arr = resolve(row, "final") or [0] * 8
        ads_arr = resolve(row, "adsBudget") or [0] * 8
        promo_arr = resolve(row, "promoBudget") or [0] * 8
        rrp = resolve(row, "rrp") or 0
        target_units = final_arr[month_index] if month_index < len(final_arr) else 0

        sku_updates.append({
            "sku": sku,
            "channel": resolve(row, "channel") or "N/A",
            "portfolio": resolve(row, "portfolio") or "N/A",
            "priority": resolve(row, "priority") or "N/A",
            "moc": resolve(row, "moc") or 0,
            "moc_band": resolve(row, "mocBand") or "N/A",
            "rrp": round(rrp, 2),
            "updated_at": datetime.utcnow().isoformat(),
        })
        target_rows.append({
            "sku": sku,
            "month": month_date,
            "target_units": target_units,
            "target_gmv": round(target_units * rrp, 2),
            "ads_target": round(ads_arr[month_index] if month_index < len(ads_arr) else 0, 2),
            "promo_target": round(promo_arr[month_index] if month_index < len(promo_arr) else 0, 2),
            "updated_at": datetime.utcnow().isoformat(),
        })

    if sku_filter is not None:
        print(f"[target] scoped to {len(sku_updates)} managed SKU(s), skipped {skipped} outside the filter")

    # upsert skus first (creates a placeholder row if `followup` hasn't run yet)
    # so the targets_monthly foreign key never fails regardless of run order
    upsert_rows("skus", sku_updates, on_conflict="sku")
    upsert_rows("targets_monthly", target_rows, on_conflict="sku,month")


def ingest_target_csv(path, month_label=None, sku_filter=None):
    """Ingest target/budget data from a flat CSV export (e.g.
    'SSO_US_TOTAL_Target_BudgetV2...csv') — the same underlying planning data
    as the SSO HTML dashboard, but exported as plain columns instead of a
    compact embedded-JSON payload, so it's more robust to parse and doesn't
    break if the dashboard's internal schema changes.
    Expected columns (case-sensitive, as exported): SKU, Product Name,
    Product Line, Portfolio, MOC, MOC Band, Month, Final Units, Current RRP,
    Final GMV, Total Ads Budget, Promotion Budget, Channel."""
    with open(path, encoding="utf-8-sig", newline="") as f:
        all_rows = list(csv.DictReader(f))
    if not all_rows:
        print("[target-csv] file is empty")
        return

    months_present = sorted(set(r.get("Month", "") for r in all_rows if r.get("Month")))
    target_month_label = month_label or (months_present[0] if len(months_present) == 1 else None)
    if target_month_label is None:
        sys.exit(f"[target-csv] multiple months found in file ({months_present}) — pass --month to pick one")
    month_date = target_month_label if re.match(r"^\d{4}-\d{2}-\d{2}$", target_month_label) else _label_to_date(target_month_label)

    def num(v, default=0.0):
        if v is None or v == "":
            return default
        try:
            return float(str(v).replace(",", "").replace("%", ""))
        except (TypeError, ValueError):
            return default

    sku_updates, target_rows = [], []
    skipped = 0
    for row in all_rows:
        if row.get("Month") != target_month_label:
            continue
        sku = (row.get("SKU") or "").strip()
        if not sku:
            continue
        if sku_filter is not None and sku not in sku_filter:
            skipped += 1
            continue
        rrp = num(row.get("Current RRP"))
        sku_updates.append({
            "sku": sku,
            "channel": row.get("Channel") or "N/A",
            "portfolio": row.get("Portfolio") or "N/A",
            "priority": "N/A",  # not present in this CSV export
            "moc": num(row.get("MOC")),
            "moc_band": row.get("MOC Band") or "N/A",
            "rrp": round(rrp, 2),
            "updated_at": datetime.utcnow().isoformat(),
        })
        target_rows.append({
            "sku": sku,
            "month": month_date,
            "target_units": num(row.get("Final Units")),
            "target_gmv": round(num(row.get("Final GMV")), 2),
            "ads_target": round(num(row.get("Total Ads Budget")), 2),
            "promo_target": round(num(row.get("Promotion Budget")), 2),
            "updated_at": datetime.utcnow().isoformat(),
        })

    if sku_filter is not None:
        print(f"[target-csv] scoped to {len(sku_updates)} managed SKU(s), skipped {skipped} row(s) outside the filter")

    upsert_rows("skus", sku_updates, on_conflict="sku")
    upsert_rows("targets_monthly", target_rows, on_conflict="sku,month")


# ---------------------------------------------------------------------------
# 3) Daily/hourly sales extract -> sales_daily
# ---------------------------------------------------------------------------
def ingest_sales_excel(path, sheet=None, source_label=None, sku_filter=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    hidx = {h: i for i, h in enumerate(headers)}
    date_col = "date" if "date" in hidx else "month"  # some extracts name it "month" but hold a full date

    def f(row, key, default=0.0):
        if key not in hidx:
            return default
        v = row[hidx[key]]
        try:
            return float(v) if v not in (None, "") else default
        except (TypeError, ValueError):
            return default

    rows = {}
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku = row[hidx["sku"]]
        if not sku:
            continue
        sku = str(sku).strip()
        if sku_filter is not None and sku not in sku_filter:
            skipped += 1
            continue
        d = row[hidx[date_col]]
        date_str = d.strftime("%Y-%m-%d") if isinstance(d, datetime) else str(d)[:10]
        key = (sku, date_str)
        agg = rows.setdefault(key, {
            "sku": sku, "date": date_str, "units": 0.0, "gmv": 0.0, "ads": 0.0, "promo": 0.0,
            "ads_gmv": 0.0, "ads_units": 0.0,
            "category": row[hidx["main_category"]] if "main_category" in hidx else None,
            "source_file": source_label or os.path.basename(path),
        })
        agg["units"] += f(row, "ordered_units")
        agg["gmv"] += f(row, "ordered_gmv")
        agg["ads"] += f(row, "total_ads")
        agg["promo"] += f(row, "total_promo")
        # Ads Performance tab: ad-attributed GMV/units across the 3 ad types
        agg["ads_gmv"] += f(row, "sb_ordered_nmv") + f(row, "sd_ordered_nmv") + f(row, "sp_ordered_nmv")
        agg["ads_units"] += f(row, "sb_ordered_units") + f(row, "sd_ordered_units") + f(row, "sp_ordered_units")
    wb.close()

    out = []
    skus_seen = set()
    for r in rows.values():
        r["units"] = round(r["units"], 2)
        r["gmv"] = round(r["gmv"], 2)
        r["ads"] = round(r["ads"], 2)
        r["promo"] = round(r["promo"], 2)
        r["ads_gmv"] = round(r["ads_gmv"], 2)
        r["ads_units"] = round(r["ads_units"], 2)
        out.append(r)
        skus_seen.add(r["sku"])

    if sku_filter is not None:
        print(f"[sales] scoped to {len(skus_seen)} managed SKU(s) / {len(out)} day-rows, skipped {skipped} row(s) outside the filter")

    # make sure every SKU in this file exists in `skus` first (placeholder if
    # new), otherwise the sales_daily foreign key would reject unknown SKUs
    # regardless of whether `followup`/`target` have been run yet
    placeholders = [{"sku": s, "updated_at": datetime.utcnow().isoformat()} for s in skus_seen]
    upsert_rows("skus", placeholders, on_conflict="sku")
    upsert_rows("sales_daily", out, on_conflict="sku,date")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="Ingest Yes4All source files into Supabase")
    sub = p.add_subparsers(dest="cmd", required=True)

    p_fu = sub.add_parser("followup", help="Amazon Follow up_2026.xlsx -> skus")
    p_fu.add_argument("path")
    p_fu.add_argument("--sheet", default="Tracking (2)")

    p_tg = sub.add_parser("target", help="SSO_US planning HTML -> skus + targets_monthly")
    p_tg.add_argument("path")
    p_tg.add_argument("--month", default=None, help="YYYY-MM-DD or 'Aug-26' style label; default = first future month in the file")
    p_tg.add_argument("--month-index", type=int, default=0, help="index into the file's future-month array, default 0 (first)")
    p_tg.add_argument("--skus-from", default=None, help="path to the follow-up/tracking Excel file — if given, only SKUs found there are ingested (the SSO export covers the whole company's catalog, not just the managed SKU list)")

    p_tc = sub.add_parser("target-csv", help="Flat SSO_US_TOTAL_Target_Budget...csv export -> skus + targets_monthly (more robust than the HTML payload)")
    p_tc.add_argument("path")
    p_tc.add_argument("--month", default=None, help="e.g. 'Aug-26' or '2026-08-01'; default = auto-detect if the file only has one month")
    p_tc.add_argument("--skus-from", default=None, help="path to the follow-up/tracking Excel file — if given, only SKUs found there are ingested")

    p_sa = sub.add_parser("sales", help="SSO daily/hourly extraction Excel -> sales_daily")
    p_sa.add_argument("path")
    p_sa.add_argument("--sheet", default=None)
    p_sa.add_argument("--skus-from", default=None, help="path to the follow-up/tracking Excel file — if given, only SKUs found there are ingested")

    p_cl = sub.add_parser("cleanup", help="One-time: DELETE any skus row (and cascaded targets/sales) NOT in the follow-up file's managed SKU list")
    p_cl.add_argument("path", help="path to the follow-up/tracking Excel file (the managed SKU list)")

    p.add_argument("--out", default=None, help="output .sql path (sql mode only); default = <cmd>_upsert.sql next to this script")

    args = p.parse_args()
    if args.cmd == "followup":
        ingest_followup(args.path, sheet=args.sheet)
    elif args.cmd == "target":
        sku_filter = load_managed_skus(args.skus_from) if args.skus_from else None
        ingest_target(args.path, month_label=args.month, month_index=args.month_index, sku_filter=sku_filter)
    elif args.cmd == "target-csv":
        sku_filter = load_managed_skus(args.skus_from) if args.skus_from else None
        ingest_target_csv(args.path, month_label=args.month, sku_filter=sku_filter)
    elif args.cmd == "sales":
        sku_filter = load_managed_skus(args.skus_from) if args.skus_from else None
        ingest_sales_excel(args.path, sheet=args.sheet, sku_filter=sku_filter)
    elif args.cmd == "cleanup":
        cleanup_unmanaged_skus(args.path)

    if INGEST_MODE != "api" and SQL_BUFFER:
        out_path = args.out or os.path.join(os.path.dirname(os.path.abspath(__file__)), f"{args.cmd}_upsert.sql")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("begin;\n\n" + "\n\n".join(SQL_BUFFER) + "\n\ncommit;\n")
        print(f"\nWritten {out_path}")
        print("Paste its contents into Supabase -> SQL Editor -> New query -> Run.")


if __name__ == "__main__":
    main()
